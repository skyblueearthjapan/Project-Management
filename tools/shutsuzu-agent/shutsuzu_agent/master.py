"""工番マスタ Excel（新一覧.xlsm / 日程表A.xlsx）読込（参照EXE :26-156 流用）。

新一覧・日程表A は **社内全体で使う共有マスタ**のため、EXE フォルダには置かず
ファイルサーバの固定パス（config の ``master_nittei_path`` / ``master_shin_ichiran_path``）
を参照する。**日程表A を優先**し、無い工番は **新一覧** で補完する（元アプリと同方式）。

openpyxl は ``read_only`` / ``data_only`` で読むだけ。書込・削除・改名は一切しない。
"""

from __future__ import annotations

import re
from collections.abc import Callable
from typing import Any

import openpyxl


def _load_rows(
    path: str, sheet_filter: Callable[[str], bool] | None = None
) -> dict[str, dict[str, str]]:
    """Excel から ``{工番: {nohin(納入先), hinmei(品名)}}`` を読む。

    列: 0=工番, 2=納入先(nohin), 3=品名(hinmei)（参照EXE と同じ並び）。
    """
    result: dict[str, dict[str, str]] = {}
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        for sheet_name in wb.sheetnames:
            if sheet_filter is not None and not sheet_filter(sheet_name):
                continue
            ws = wb[sheet_name]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or not row[0]:
                    continue
                no = str(row[0]).strip()
                result[no] = {
                    "nohin": (str(row[2]).strip() if len(row) > 2 and row[2] else ""),
                    "hinmei": (str(row[3]).strip() if len(row) > 3 and row[3] else ""),
                }
    finally:
        wb.close()
    return result


def load_master(nittei_path: str, shin_ichiran_path: str) -> dict[str, dict[str, str]]:
    """新一覧(ベース) → 日程表A(上書き) の順でマスタ辞書を構築する（日程表A 優先）。

    どちらか一方が読めなくても、読めた側のデータだけで継続する。
    """
    master: dict[str, dict[str, str]] = {}
    # 新一覧（シート "正"）をベースに。
    try:
        master.update(_load_rows(shin_ichiran_path, sheet_filter=lambda sn: sn == "正"))
    except (OSError, ValueError, KeyError):
        pass
    # 日程表A（"日程表" を含むシート）で上書き（優先）。
    try:
        master.update(_load_rows(nittei_path, sheet_filter=lambda sn: "日程表" in sn))
    except (OSError, ValueError, KeyError):
        pass
    return master


def _candidate_keys(job_id: str) -> list[str]:
    """job_id（``LW25146`` / ``TS26007`` 等）から Excel 照合キー候補を作る。

    Excel は LW 工番を 5桁番号（``25146``）、TS 等を prefix 込み（``TS26007``）で持つため、
    完全一致 → ``LW`` 除去 → 数字のみ の順で試す。
    """
    jid = job_id.strip()
    keys = [jid]
    if jid.upper().startswith("LW") and jid[2:]:
        keys.append(jid[2:])
    digits = re.sub(r"\D", "", jid)
    if digits and digits not in keys:
        keys.append(digits)
    return keys


def lookup(master: dict[str, dict[str, str]], job_id: str) -> dict[str, Any] | None:
    """マスタを引き ``customer(納入先) / title(品名) / delivery_date`` で返す。無ければ None。

    delivery_date は参照EXE 同様このExcelからは取得しない（None）。
    """
    for key in _candidate_keys(job_id):
        info = master.get(key)
        if info:
            return {
                "customer": info.get("nohin") or None,
                "title": info.get("hinmei") or None,
                "delivery_date": None,
            }
    return None
